"""
Extrae los referentes de presupuesto (Hospital Civil Ipiales y Casa Ricaurte) a JSON.
Uso: python extract_referentes.py <salida.json>
"""
import openpyxl, json, sys, re, datetime

BASE = r"D:\DPTO. ARQUITECTURA\A - Ing. Ana\CASA RICAURTE"
F_HOSP = BASE + r"\PRESUPUESTO_H CIVIL IPIALES_OBRACIVIL_2026_07_28_REFERENTE.xlsx"
F_RIC = BASE + r"\PRESUPUESTO_CANTIDADES_CASA_RICAURTE.xlsx"


def rows_of(wb, name):
    out = []
    for r in wb[name].iter_rows(values_only=True):
        v = list(r)
        while v and v[-1] is None:
            v.pop()
        out.append(v)
    return out


def g(r, i):
    return r[i] if r and len(r) > i else None


def num(x):
    try:
        if x is None or x == "":
            return None
        return float(x)
    except Exception:
        return None


def item_code(x):
    """'2001-02-01' (item 1.2.1 leído como fecha) -> '1.2.1'"""
    if isinstance(x, datetime.datetime):
        return f"{x.day}.{x.month}.{x.year - 2000}"
    if isinstance(x, (int, float)):
        return str(x)
    s = str(x or "").strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{int(m.group(3))}.{int(m.group(2))}.{int(m.group(1)) - 2000}"
    return s


BUCKETS = {5: "equipo", 6: "material", 7: "mano_obra", 8: "transporte"}


def parse_apu_blocks(rows, fuente):
    """Bloques: [item, descripcion, ..., unidad(idx8)] / encabezados / líneas / C. DIRECTO"""
    apus = []
    i = 0
    n = len(rows)
    while i < n:
        r = rows[i]
        nxt = rows[i + 1] if i + 1 < n else []
        if g(r, 1) and g(r, 8) and str(g(nxt, 5) or "").upper().startswith("EQUIPO"):
            apu = {"fuente": fuente, "item": item_code(g(r, 0)), "descripcion": str(g(r, 1)).strip(),
                   "unidad": str(g(r, 8)).strip().upper(), "lineas": []}
            j = i + 3
            while j < n:
                l = rows[j]
                d = g(l, 1)
                if d and str(d).strip().upper().startswith("C. DIRECTO"):
                    apu["directo"] = num(g(l, 4))
                    apu["subtotales"] = {"equipo": num(g(l, 5)) or 0, "material": num(g(l, 6)) or 0,
                                         "mano_obra": num(g(l, 7)) or 0, "transporte": num(g(l, 8)) or 0}
                    break
                if d and num(g(l, 3)) is not None:
                    bucket = None
                    for k in (5, 6, 7, 8):
                        if num(g(l, k)):
                            bucket = BUCKETS[k]
                    du = str(d).upper()
                    if bucket is None:
                        bucket = "mano_obra" if "CUADRILLA" in du else "material"
                    apu["lineas"].append({"descripcion": str(d).strip(), "unidad": str(g(l, 2) or "").strip().upper(),
                                          "cantidad": num(g(l, 3)), "precio": num(g(l, 4)) or 0, "bucket": bucket})
                j += 1
            apus.append(apu)
            i = j + 1
        else:
            i += 1
    return apus


def parse_basicos(rows, fuente):
    out = []
    i = 0
    while i < len(rows):
        r = rows[i]
        if g(r, 0) and str(g(r, 1) or "").upper() == "UNIDAD" and str(g(r, 2) or "").upper().startswith("RENDIM"):
            b = {"fuente": fuente, "descripcion": str(g(r, 0)).strip(), "lineas": []}
            j = i + 1
            while j < len(rows):
                l = rows[j]
                if str(g(l, 0) or "").upper().startswith("DIRECTO"):
                    b["unidad"] = str(g(l, 1) or "M3").upper()
                    b["directo"] = num(g(l, 3))
                    break
                if g(l, 0) and num(g(l, 2)) is not None:
                    bucket = "material"
                    for k, name in ((4, "equipo"), (5, "material"), (6, "mano_obra"), (7, "transporte")):
                        if num(g(l, k)):
                            bucket = name
                    b["lineas"].append({"descripcion": str(g(l, 0)).strip(), "unidad": str(g(l, 1) or "").upper(),
                                        "cantidad": num(g(l, 2)), "precio": num(g(l, 3)) or 0, "bucket": bucket})
                j += 1
            out.append(b)
            i = j + 1
        else:
            i += 1
    return out


def parse_materiales(rows, fuente, has_code):
    out = []
    for r in rows[3:]:
        desc = g(r, 2)
        if not desc or num(g(r, 5)) is None:
            continue
        cots = []
        for base in (6, 9, 12):
            prov, val = g(r, base + 1), num(g(r, base + 2))
            if prov and val:
                cots.append({"proveedor": str(prov).strip().upper(), "valor": val, "referencia": str(g(r, base) or "")[:300]})
        out.append({"fuente": fuente, "categoria": str(g(r, 1) or "").strip().upper(), "descripcion": str(desc).strip(),
                    "codigo_erp": (str(g(r, 3)).strip() if has_code and g(r, 3) else None),
                    "unidad": str(g(r, 4) or "UND").strip().upper(), "precio_ref": num(g(r, 5)), "cotizaciones": cots})
    return out


def parse_cuadro(rows, fuente, desc_col=1, unit_col=2, qty_col=3, price_col=4):
    caps, cur, sub = [], None, None
    for r in rows:
        it, d = g(r, 0), g(r, desc_col)
        if not d:
            continue
        u, q = g(r, unit_col), num(g(r, qty_col))
        code = item_code(it) if it is not None else ""
        if u and q:
            if cur is None:
                cur = {"capitulo": "GENERAL", "items": []}
                caps.append(cur)
            cur["items"].append({"item": code, "descripcion": str(d).strip(), "unidad": str(u).strip().upper(),
                                 "cantidad": q, "precio": num(g(r, price_col)), "subcapitulo": sub,
                                 "codigo_erp": (str(g(r, 6)).strip() if fuente == "RICAURTE" and g(r, 6) else None)})
        elif it is not None and not str(d).upper().startswith("SUBTOTAL"):
            depth = code.count(".")
            if fuente == "HOSPITAL" and depth == 0:
                continue  # 1 = OBRA CIVIL (macro)
            if fuente == "HOSPITAL" or code.endswith(".00") or depth <= 1:
                cur = {"capitulo": str(d).strip(), "item": code, "items": []}
                caps.append(cur)
                sub = None
    return [c for c in caps if c["items"]]


def main(out_path):
    data = {"salarios": [], "cuadrillas": [], "factor_prestacional": {}, "materiales": [], "equipos": [],
            "transportes": [], "apus": [], "basicos": [], "cuadros": {}, "apu_erp_ricaurte": []}

    wb = openpyxl.load_workbook(F_HOSP, read_only=True, data_only=True)
    mo = rows_of(wb, "MANO DE OBRA")
    for r in mo[22:48]:
        if g(r, 1) and num(g(r, 2)) and num(g(r, 3)):
            data["salarios"].append({"cargo": str(g(r, 1)).strip(), "salario": num(g(r, 2)),
                                     "factor": num(g(r, 3)), "mensual": num(g(r, 4)), "jornal": num(g(r, 5))})
    i = 60
    while i < len(mo):
        r = mo[i]
        if g(r, 1) and str(g(r, 1)).upper().startswith("CUADRILLA") and str(g(r, 3) or "").upper() == "H-H":
            c = {"nombre": str(g(r, 1)).strip(), "integrantes": []}
            j = i + 2
            while j < len(mo) and not str(g(mo[j], 1) or "").upper().startswith("TOTAL"):
                if g(mo[j], 1) and num(g(mo[j], 2)):
                    c["integrantes"].append({"cargo": str(g(mo[j], 1)).strip(), "cantidad": num(g(mo[j], 2))})
                j += 1
            if j + 1 < len(mo):
                c["valor_hora"] = num(g(mo[j + 1], 3))
            data["cuadrillas"].append(c)
            i = j
        i += 1

    fp = rows_of(wb, "F PRESTACIONAL")
    def fp_block(start, end, cols):
        comps = []
        for r in fp[start:end]:
            if g(r, 0) and num(g(r, cols[0])) is not None and g(r, 0) not in ("Mensual", "Anual"):
                comps.append({"concepto": str(g(r, 0)).strip(), "pct": [num(g(r, c)) for c in cols]})
        return comps
    data["factor_prestacional"] = {
        "smlv": num(g(fp[7], 3)), "subsidio_transporte": num(g(fp[4], 2)),
        "obreros": {"niveles": [1, 1.5, 2], "componentes": fp_block(13, 32, (4, 6, 8)),
                    "factor": [num(g(fp[32], 4)), num(g(fp[32], 6)), num(g(fp[32], 8))]},
        "empleados": {"niveles": [2.5, 3, 4], "componentes": fp_block(49, 68, (4, 6, 8)),
                      "factor": [num(g(fp[68], 4)), num(g(fp[68], 6)), num(g(fp[68], 8))]},
    }
    data["materiales"] += parse_materiales(rows_of(wb, "MATERIALES"), "HOSPITAL", False)
    for r in rows_of(wb, "EQUIPOS")[1:]:
        if g(r, 1) and num(g(r, 3)):
            data["equipos"].append({"descripcion": str(g(r, 1)).strip(), "unidad": str(g(r, 2) or "HORA").upper(), "precio": num(g(r, 3))})
    for r in rows_of(wb, "TRASPORTES")[1:4]:
        if g(r, 1) and num(g(r, 3)):
            data["transportes"].append({"descripcion": str(g(r, 1)).strip(), "unidad": str(g(r, 2)).upper(), "precio": num(g(r, 3))})
    data["apus"] += parse_apu_blocks(rows_of(wb, "APU O CIVL"), "HOSPITAL")
    data["apus"] += parse_apu_blocks(rows_of(wb, "APU EXTINCION"), "HOSPITAL_RCI")
    data["basicos"] += parse_basicos(rows_of(wb, "APU BASE"), "HOSPITAL")
    data["cuadros"]["HOSPITAL"] = parse_cuadro(rows_of(wb, "CUADRO PRESUPUESTO"), "HOSPITAL")

    wb = openpyxl.load_workbook(F_RIC, read_only=True, data_only=True)
    data["materiales"] += parse_materiales(rows_of(wb, "MATERIALES"), "RICAURTE", True)
    data["apus"] += parse_apu_blocks(rows_of(wb, "APU O CIVL"), "RICAURTE")
    data["basicos"] += parse_basicos(rows_of(wb, "APU BASE"), "RICAURTE")
    data["cuadros"]["RICAURTE"] = parse_cuadro(rows_of(wb, "CUADRO PRESUPUESTO"), "RICAURTE")
    for r in rows_of(wb, "APU_ERP")[1:]:
        if g(r, 1) and g(r, 6):
            data["apu_erp_ricaurte"].append({"tipo": str(g(r, 0) or "").upper(), "descripcion": str(g(r, 1)).strip(),
                                             "precio": num(g(r, 2)) or 0, "unidad": str(g(r, 3) or "").upper(),
                                             "cantidad": num(g(r, 4)) or 0, "item": str(g(r, 5) or ""),
                                             "codigo_apu": str(g(r, 6)).strip().upper(),
                                             "codigo_insumo": str(g(r, 7) or "").strip().upper()})

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1, default=str)
    resumen = {k: len(v) for k, v in data.items() if isinstance(v, list)}
    resumen["cuadros"] = {k: sum(len(c["items"]) for c in v) for k, v in data["cuadros"].items()}
    resumen["fp_obreros"] = data["factor_prestacional"]["obreros"]["factor"]
    print(json.dumps(resumen))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "referentes.json")
